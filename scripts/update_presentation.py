#!/usr/bin/env python3
"""Update VIK presentation content while preserving formatting and template."""
from pptx import Presentation
from pptx.util import Pt
from copy import deepcopy
from lxml import etree
import os

INPUT = "docs/Lazarev-Presentation-VIK-25-26.pptx"
OUTPUT = "docs/Lazarev-Presentation-VIK-25-26-updated.pptx"

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def clear_and_set(shape, new_text):
    """Clear ALL text from shape and set new multi-paragraph text.
    Preserves formatting (font, size, color) from the first run."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    txBody = tf._txBody

    # Save first run formatting
    first_rPr = None
    first_pPr = None
    for p in txBody.findall('.//a:p', NS):
        pPr = p.find('a:pPr', NS)
        if pPr is not None and first_pPr is None:
            first_pPr = deepcopy(pPr)
        for r in p.findall('a:r', NS):
            rPr = r.find('a:rPr', NS)
            if rPr is not None and first_rPr is None:
                first_rPr = deepcopy(rPr)
            break
        if first_rPr is not None:
            break

    # Remove ALL existing paragraphs
    for p in txBody.findall('a:p', NS):
        txBody.remove(p)

    # Add new paragraphs
    lines = new_text.split('\n')
    aNS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    for line in lines:
        p = etree.SubElement(txBody, f'{{{aNS}}}p')
        if first_pPr is not None:
            p.append(deepcopy(first_pPr))
        if line.strip():  # non-empty line
            r = etree.SubElement(p, f'{{{aNS}}}r')
            if first_rPr is not None:
                r.append(deepcopy(first_rPr))
            else:
                rPr = etree.SubElement(r, f'{{{aNS}}}rPr')
                rPr.set('lang', 'ru-RU')
                rPr.set('dirty', '0')
            t = etree.SubElement(r, f'{{{aNS}}}t')
            t.text = line


def main():
    prs = Presentation(INPUT)
    slides = list(prs.slides)

    # ===== SLIDE 2: Актуальность и результаты =====
    slide2 = slides[1]
    for s in slide2.shapes:
        if not s.has_text_frame:
            continue
        text = s.text_frame.text
        if "Актуальность тематики" in text:
            clear_and_set(s, "Актуальность исследования")
        elif "Ожидаемые результаты" in text:
            clear_and_set(s, "Результаты исследования")
        elif "Временные ряды лежат" in text:
            clear_and_set(s,
                "Временные ряды являются основным инструментом моделирования "
                "процессов в экономике, энергетике, транспорте и финансах. "
                "Классические статистические модели (ARIMA, ETS) требуют "
                "ручной настройки и теряют качество при нарушении стационарности.\n"
                "\n"
                "Современные нейросетевые архитектуры на основе трансформеров "
                "(Autoformer, FEDformer, PatchTST) обещают лучше захватывать "
                "долгосрочные зависимости, однако их эффективность критически "
                "зависит от способа декомпозиции входного ряда.\n"
                "\n"
                "Актуален вопрос: при каких условиях обучаемая декомпозиция "
                "и нейросетевые механизмы действительно улучшают прогноз, "
                "а при каких более простые подходы предпочтительнее?"
            )
        elif "Реализовать и исследовать" in text:
            clear_and_set(s,
                "Реализован экспериментальный контур из 11 моделей: "
                "4 классические (ARIMA, ETS, Prophet, Seasonal Naive), "
                "6 нейросетевых (DLinear, N-BEATS, TimesNet, Autoformer, "
                "FEDformer, PatchTST) и гибридная Helformer.\n"
                "\n"
                "Сравнение проведено на 180 рядах из M3 и M4 Competition "
                "по 6 метрикам (sMAPE, MAPE, RMSE, MSE, MAE, MASE).\n"
                "\n"
                "Исследованы три режима оценки: прямой прогноз, "
                "итеративное прогнозирование и предобучение трансформеров.\n"
                "\n"
                "Helformer (фиксированная декомпозиция + нейросетевая "
                "коррекция) показал лучший sMAPE 10,61% на M3."
            )

    # ===== SLIDE 3: Новизна и значимость =====
    slide3 = slides[2]
    for s in slide3.shapes:
        if not s.has_text_frame:
            continue
        text = s.text_frame.text
        if "Новизна планируемых" in text:
            clear_and_set(s, "Научная новизна")
        elif "Практическая значимость планируемых" in text:
            clear_and_set(s, "Практическая значимость")
        elif "Выполняется углублённое" in text:
            clear_and_set(s,
                "Проведено систематическое сравнение фиксированной "
                "и обучаемой декомпозиции временных рядов на 180 рядах "
                "из M3 и M4 Competition с участием 11 моделей.\n"
                "\n"
                "Установлено, что гибридный подход Helformer "
                "(фиксированная декомпозиция Хольта\u2013Уинтерса "
                "+ нейросетевая коррекция) превосходит как чисто "
                "статистические, так и чисто нейросетевые методы.\n"
                "\n"
                "Выявлена граница ~50\u2013100 наблюдений, ниже которой "
                "нейросети уступают классическим моделям.\n"
                "\n"
                "Показано, что предобучение трансформеров на корпусе "
                "рядов даёт улучшение на 7\u201311% sMAPE."
            )
        elif "Улучшение качества" in text:
            clear_and_set(s,
                "Разработан воспроизводимый программный стенд на Python "
                "(PyTorch), позволяющий сравнивать произвольные модели "
                "на стандартных бенчмарках с единым протоколом.\n"
                "\n"
                "Результаты применимы для выбора оптимальной стратегии "
                "прогнозирования в зависимости от длины ряда и типа данных.\n"
                "\n"
                "Рекомендации: для коротких рядов с сезонностью оптимален "
                "гибридный подход Helformer, для длинных \u2014 DLinear "
                "или N-BEATS."
            )

    # ===== SLIDE 4: Описание проекта =====
    slide4 = slides[3]
    for s in slide4.shapes:
        if not s.has_text_frame:
            continue
        if "Проект посвящён" in s.text_frame.text:
            clear_and_set(s,
                "Исследование посвящено сравнительному анализу подходов "
                "к прогнозированию одномерных временных рядов с акцентом "
                "на роли декомпозиции в нейросетевых архитектурах.\n"
                "\n"
                "Центральный вопрос: в каких условиях обучаемая декомпозиция "
                "(DLinear, Autoformer, N-BEATS) эффективнее фиксированной "
                "(ETS, Helformer), и наоборот.\n"
                "\n"
                "Экспериментальная база: 90 рядов M3 (годовые, квартальные, "
                "месячные) и 90 рядов M4 (квартальные, месячные, ежедневные) "
                "\u2014 всего 180 рядов с горизонтами от 6 до 18 шагов.\n"
                "\n"
                "Ключевые результаты:\n"
                "\u2022 Helformer \u2014 лучшая модель на M3 "
                "(sMAPE 10,61%, MASE 0,94)\n"
                "\u2022 ETS \u2014 лучшая среди статистических "
                "(sMAPE 13,89%)\n"
                "\u2022 DLinear \u2014 лучший на ежедневных рядах M4 "
                "(sMAPE 2,43%)\n"
                "\u2022 Предобучение улучшает трансформеры на 7\u201311%\n"
                "\u2022 Autoformer и FEDformer наименее устойчивы при "
                "итеративном прогнозе"
            )

    # ===== SLIDE 5: План-график — оставляем таблицу как есть =====

    # ===== SLIDE 6: Заключение и выводы =====
    slide6 = slides[5]
    for s in slide6.shapes:
        if not s.has_text_frame:
            continue
        if "В заключении предполагается" in s.text_frame.text:
            clear_and_set(s,
                "На основании проведённых экспериментов сформулированы "
                "следующие выводы.\n"
                "\n"
                "1. Helformer (фиксированная декомпозиция Хольта\u2013Уинтерса "
                "+ трансформерная коррекция) \u2014 лучшая модель на M3 "
                "(sMAPE 10,61%, средний ранг 2,00 из 11).\n"
                "\n"
                "2. Длина ряда \u2014 критический фактор: при менее "
                "50 наблюдениях нейросети уступают ETS и ARIMA; "
                "при более 100 \u2014 конкурируют.\n"
                "\n"
                "3. Обучаемая декомпозиция (DLinear, N-BEATS) эффективна "
                "на длинных рядах: DLinear \u2014 лучший на M4 daily "
                "(sMAPE 2,43%).\n"
                "\n"
                "4. Предобучение на корпусе рядов улучшает трансформеры "
                "на 7\u201311%, делая их конкурентоспособными.\n"
                "\n"
                "5. При итеративном прогнозировании модели с обучаемой "
                "декомпозицией деградируют сильнее всего (FEDformer "
                "+5 п.п.), PatchTST \u2014 наиболее устойчивый (+2 п.п.)."
            )

    prs.save(OUTPUT)
    print(f"Saved to {OUTPUT}")


if __name__ == "__main__":
    os.chdir("/Applications/forvscode/BachelorFinal/Hybrid-models-for-TS-processing")
    main()
